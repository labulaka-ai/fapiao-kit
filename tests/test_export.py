"""Excel：用 openpyxl 回读，校验合计 = 各行之和、大写金额、发票号按文本存、类别小计。"""
from decimal import Decimal

import openpyxl
import pytest

import fapiao_kit
from fapiao_kit.money import to_chinese_upper


@pytest.fixture(scope="module")
def wb(sample_set, tmp_path_factory):
    d, truth = sample_set
    path = tmp_path_factory.mktemp("x") / "汇总.xlsx"
    fapiao_kit.export_xlsx(fapiao_kit.scan(d, truth["rules"]), str(path))
    return openpyxl.load_workbook(path), truth


def _detail(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    body = [r for r in rows[1:] if isinstance(r[0], int)]
    tail = {r[0]: r for r in rows[1:] if isinstance(r[0], str)}
    return header, body, tail


def test_sheets(wb):
    book, _ = wb
    assert book.sheetnames == ["汇总", "按类别小计", "需处理"]


def test_rows_and_totals(wb):
    book, truth = wb
    header, body, tail = _detail(book["汇总"])
    col = {h: i for i, h in enumerate(header)}
    assert len(body) == len(truth["records"])
    exp_total = sum(Decimal(r["fields"]["total"]) for r in truth["records"].values())
    for k in ("金额", "税额", "价税合计"):
        s = sum(Decimal(str(r[col[k]])) for r in body)
        assert Decimal(str(tail["合计"][col[k]])).quantize(Decimal("0.01")) == s.quantize(Decimal("0.01"))
    assert Decimal(str(tail["合计"][col["价税合计"]])).quantize(Decimal("0.01")) == exp_total
    assert tail["合计（大写）"][2] == to_chinese_upper(exp_total)
    assert tail["张数"][2] == len(truth["records"])
    assert "未识别" in tail


def test_invoice_number_is_text(wb):
    book, truth = wb
    ws = book["汇总"]
    header = [c.value for c in ws[1]]
    ci = header.index("发票号码") + 1
    nums = []
    for row in ws.iter_rows(min_row=2):
        c = row[ci - 1]
        if isinstance(row[0].value, int):
            assert c.data_type == "s" and c.number_format == "@"
            nums.append(c.value)
    assert sorted(nums) == sorted(r["fields"]["number"] for r in truth["records"].values())
    assert "99260000000000000001" in nums  # 20 位号码没有变成科学计数法


def test_category_subtotals(wb):
    book, truth = wb
    rows = list(book["按类别小计"].iter_rows(values_only=True))
    cats = {r[0]: r for r in rows[1:-1]}
    exp = {}
    for r in truth["records"].values():
        c = r["fields"]["category"]
        exp[c] = exp.get(c, Decimal(0)) + Decimal(r["fields"]["total"])
    assert set(cats) == set(exp)
    for c, v in exp.items():
        assert Decimal(str(cats[c][4])).quantize(Decimal("0.01")) == v
    assert rows[-1][0] == "合计" and rows[-1][1] == len(truth["records"])


def test_todo_sheet_lists_problems(wb):
    book, truth = wb
    rows = list(book["需处理"].iter_rows(values_only=True))[1:]
    text = "\n".join(str(r) for r in rows)
    assert "示例大字" in text and "99TEST000000000009" in text and "价税合计" in text
    for n in truth["unrecognized"]:
        assert n in text


def test_empty_export(tmp_path):
    p = tmp_path / "e.xlsx"
    fapiao_kit.export_xlsx([], str(p))
    ws = openpyxl.load_workbook(p)["汇总"]
    assert ws["A2"].value == "合计"


def test_libreoffice_can_open_xlsx(wb, sample_set, tmp_path):
    """可选：用 LibreOffice 打开导出的 xlsx 并转成 CSV，确认合计行可读（未安装 LibreOffice 时跳过）。"""
    import shutil
    import subprocess
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("未安装 LibreOffice")
    d, truth = sample_set
    x = tmp_path / "lo.xlsx"
    fapiao_kit.export_xlsx(fapiao_kit.scan(d, truth["rules"]), str(x))
    subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp_path}/lo-profile",
                    "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76", "--outdir", str(tmp_path), str(x)],
                   check=True, capture_output=True, timeout=180)
    csv = (tmp_path / "lo.csv").read_text(encoding="utf-8")
    exp_total = sum(Decimal(r["fields"]["total"]) for r in truth["records"].values())
    assert "99260000000000000001" in csv
    assert to_chinese_upper(exp_total) in csv
    assert f"{exp_total:,.2f}" in csv or f"{exp_total:.2f}" in csv
